#!/usr/bin/env python3
"""build_pptx.py – baut NOMAD-Praesentation.pptx aus folien.py.
Aufruf: ~/.venvs/pptx/bin/python build_pptx.py [Ausgabedatei]
Bilder: bilder/<name>.png (ComfyUI, comfy_bilder.py) oder Screenshots aus handbuch/bilder/annot|roh bzw. admin/public/docs."""
import os, re, sys
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, '..', '..'))
sys.path.insert(0, HERE)
from folien import SLIDES

BRAUN, HELL, AMBER, ROT = RGBColor(0x3B, 0x40, 0x21), RGBColor(0xF4, 0xED, 0xD8), RGBColor(0xD9, 0x8E, 0x04), RGBColor(0x9B, 0x1C, 0x1C)
TEXT, GRAU = RGBColor(0x1D, 0x1D, 0x1D), RGBColor(0x55, 0x55, 0x55)
FONT = 'Arial'
W, H = 13.333, 7.5
CACHE = os.path.join(HERE, 'bilder', '_cache')
os.makedirs(CACHE, exist_ok=True)


def bildpfad(name):
    for p in (f'bilder/{name}.jpg', f'bilder/{name}.png', f'../handbuch/bilder/annot/{name}.png', f'../handbuch/bilder/roh/{name}.png'):
        if os.path.exists(os.path.join(HERE, p)):
            return os.path.join(HERE, p)
    webp = os.path.join(REPO, 'admin', 'public', 'docs', name + '.webp')
    if os.path.exists(webp):
        out = os.path.join(CACHE, name + '.png')
        Image.open(webp).convert('RGB').save(out)
        return out
    raise FileNotFoundError(name)


def runs(par, text, size, farbe=TEXT, fett=False):
    for teil in re.split(r'(\*\*[^*]+\*\*)', text):
        if not teil:
            continue
        b = teil.startswith('**')
        r = par.add_run()
        r.text = teil.strip('*') if b else teil
        r.font.size, r.font.name, r.font.bold = Pt(size), FONT, (b or fett)
        r.font.color.rgb = farbe


def textbox(sl, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    return tb, tf


def balken(sl, farbe, x=0.0, y=0.0, w=0.28, h=H):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = farbe; s.line.fill.background()
    return s


def titel(sl, text, farbe=BRAUN):
    balken(sl, farbe)
    tb, tf = textbox(sl, 0.7, 0.35, 12.0, 1.0, MSO_ANCHOR.MIDDLE)
    runs(tf.paragraphs[0], text, 34, farbe, True)


def fuss(sl, nr):
    tb, tf = textbox(sl, 0.7, 7.05, 9, 0.3)
    runs(tf.paragraphs[0], 'NOMAD – deutsche Fassung (inoffiziell, Apache-2.0)', 11, GRAU)
    tb, tf = textbox(sl, 11.9, 7.05, 1.0, 0.3)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    runs(tf.paragraphs[0], str(nr), 11, GRAU)


def bild_einfuegen(sl, pfad, x, y, maxw, maxh, mitte=True, rahmen=False):
    w0, h0 = Image.open(pfad).size
    s = min(maxw / w0, maxh / h0)
    w, h = w0 * s, h0 * s
    if mitte:
        x, y = x + (maxw - w) / 2, y + (maxh - h) / 2
    pic = sl.shapes.add_picture(pfad, Inches(x), Inches(y), Inches(w), Inches(h))
    if rahmen:
        pic.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB); pic.line.width = Pt(1)
    return pic


def punkte(tf, items, size=24, farbe=TEXT, abstand=10):
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(abstand)
        runs(p, '•  ' + it, size, farbe)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
leer = prs.slide_layouts[6]

for nr, d in enumerate(SLIDES, 1):
    sl = prs.slides.add_slide(leer)
    typ = d['typ']
    if typ == 'titel':
        bild_einfuegen(sl, bildpfad(d['bild']), 0, 0, W, H, mitte=False)  # Vollbild (Seitenverhältnis 16:9 ± wenige Prozent)
        pic = sl.shapes[-1]
        pic.left, pic.top, pic.width, pic.height = 0, 0, Inches(W), Inches(H)
        ov = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.1), Inches(W), Inches(3.3))
        ov.fill.solid(); ov.fill.fore_color.rgb = BRAUN
        ov.fill.transparency = 0.15
        from pptx.oxml.ns import qn
        sf = ov.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        a = sf.makeelement(qn('a:alpha'), {'val': '82000'}); sf.append(a)
        ov.line.fill.background()
        tb, tf = textbox(sl, 0.8, 2.25, 11.7, 1.4)
        runs(tf.paragraphs[0], d['titel'], 72, RGBColor(255, 255, 255), True)
        tb, tf = textbox(sl, 0.8, 3.65, 11.7, 0.9)
        runs(tf.paragraphs[0], d['untertitel'], 34, HELL)
        tb, tf = textbox(sl, 0.8, 4.65, 11.7, 0.6)
        runs(tf.paragraphs[0], d['zeile'], 16, HELL)
    elif typ in ('bild', 'warn'):
        farbe = ROT if typ == 'warn' else BRAUN
        titel(sl, d['titel'], farbe)
        bild_einfuegen(sl, bildpfad(d['bild']), 0.7, 1.65, 6.8, 5.0, rahmen=False)
        if typ == 'warn':
            box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.75), Inches(1.6), Inches(5.2), Inches(5.2))
            box.adjustments[0] = 0.04
            box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFB, 0xEE, 0xEE); box.line.color.rgb = ROT; box.line.width = Pt(2)
        tb, tf = textbox(sl, 7.9, 1.75, 4.9, 4.9, MSO_ANCHOR.MIDDLE)
        punkte(tf, d['punkte'], 21 if len(d['punkte']) > 4 else 23)
    elif typ == 'text':
        titel(sl, d['titel'])
        if 'karten' in d:
            for i, (kt, pts) in enumerate(d['karten']):
                x = 0.7 + i * 4.15
                box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.9), Inches(3.9), Inches(4.2))
                box.adjustments[0] = 0.05
                box.fill.solid(); box.fill.fore_color.rgb = HELL; box.line.color.rgb = BRAUN; box.line.width = Pt(2)
                tb, tf = textbox(sl, x + 0.2, 2.1, 3.5, 0.8)
                runs(tf.paragraphs[0], kt, 28, BRAUN, True)
                tb, tf = textbox(sl, x + 0.2, 3.0, 3.5, 3.0)
                punkte(tf, pts, 20)
        else:
            tb, tf = textbox(sl, 0.9, 1.7, 11.6, 5.0, MSO_ANCHOR.MIDDLE)
            punkte(tf, d['punkte'], 26 if len(d['punkte']) <= 4 else 24, abstand=14)
    elif typ == 'shot':
        titel(sl, d['titel'])
        bild_einfuegen(sl, bildpfad(d['bild']), 0.9, 1.4, 11.5, 5.2, rahmen=True)
        tb, tf = textbox(sl, 0.9, 6.55, 11.5, 0.5)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        runs(tf.paragraphs[0], d['unter'], 18, GRAU)
    elif typ == 'stufen':
        titel(sl, d['titel'])
        zeilen = len(d['zeilen']) + 1
        gt = sl.shapes.add_table(zeilen, 3, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.75 * zeilen))
        t = gt.table
        t.columns[0].width, t.columns[1].width, t.columns[2].width = Inches(3.9), Inches(4.3), Inches(3.3)
        for c, ktext in enumerate(d['kopf']):
            cell = t.cell(0, c); cell.fill.solid(); cell.fill.fore_color.rgb = BRAUN
            cell.text_frame.paragraphs[0].text = ''
            runs(cell.text_frame.paragraphs[0], ktext, 20, RGBColor(255, 255, 255), True)
        for r, zeile in enumerate(d['zeilen'], 1):
            for c, ztext in enumerate(zeile):
                cell = t.cell(r, c); cell.fill.solid(); cell.fill.fore_color.rgb = HELL if r % 2 else RGBColor(255, 255, 255)
                cell.text_frame.paragraphs[0].text = ''
                runs(cell.text_frame.paragraphs[0], ztext, 19, TEXT, c == 0)
        if d.get('fuss'):
            tb, tf = textbox(sl, 0.9, 1.7 + 0.75 * zeilen + 0.3, 11.5, 0.5)
            runs(tf.paragraphs[0], d['fuss'], 15, GRAU)
    fuss(sl, nr)
    sl.notes_slide.notes_text_frame.text = d['notiz']

out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'NOMAD-Praesentation.pptx')
prs.save(out)
print(out, len(SLIDES), 'Folien')
